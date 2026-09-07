"""Write an edited text workbook back into the site.

    python3 tools/apply_strings.py editing/CAD_Explorer_text.xlsx

Re-extracts the strings first, so the spans are always current. Every edit is checked
against the file before anything is written: if the "Current text" column no longer
matches what is in the source, the run stops and names the rows rather than guessing.
index.html and app.js are patched by character span, highest offset first, so earlier
edits cannot shift later ones. data.js is rewritten from parsed JSON, and the matching
literal in build_explorer_data.py is updated too so a rebuild does not undo the change.
"""
import json, os, re, sys, subprocess
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
P = lambda *a: os.path.join(ROOT, *a)

def read_edits(xlsx):
    wb = load_workbook(xlsx)
    out = {}
    for name in wb.sheetnames:
        ws = wb[name]
        head = None
        for row in ws.iter_rows(values_only=True):
            if row and row[0] == 'ID': head = row; continue
            if not head or not row or not row[0]: continue
            rid = str(row[0]).strip()
            if not re.fullmatch(r'[HDA]\d{3}', rid): continue
            cur = (row[3] or '') if len(row) > 3 else ''
            new = (row[4] or '') if len(row) > 4 else ''
            new = str(new).strip()
            if new: out[rid] = dict(sheet_current=str(cur).strip(), new=new)
    return out

def esc_for(style, text):
    """Escape replacement text for the JS literal it is going back into."""
    if style == 'quote':    return text.replace('\\', '\\\\').replace("'", "\\'")
    if style == 'template': return text.replace('\\', '\\\\').replace('`', '\\`').replace('${', '\\${')
    return text

def main(xlsx):
    subprocess.run([sys.executable, P('tools', 'extract_strings.py')], check=True,
                   stdout=subprocess.DEVNULL)
    manifest = {r['id']: r for r in json.load(open(P('editing', 'strings.json'), encoding='utf-8'))}
    edits = read_edits(xlsx)
    if not edits:
        print('No edits found (the NEW TEXT column is empty everywhere).'); return 0

    drift = [i for i, e in edits.items()
             if i not in manifest or manifest[i]['current'] != e['sheet_current']]
    if drift:
        print('STOPPED. These rows no longer match the site, so applying them could corrupt it:')
        for i in drift:
            print(f'  {i}: sheet has {edits[i]["sheet_current"][:70]!r}')
            if i in manifest: print(f'        site has {manifest[i]["current"][:70]!r}')
        return 1

    counts = {'index.html': 0, 'data.js': 0, 'app.js': 0}

    # ---- index.html and app.js: splice by span, highest offset first ----
    for fname, pre in (('index.html', 'H'), ('app.js', 'A')):
        src = open(P(fname), encoding='utf-8').read()
        todo = sorted([manifest[i] | edits[i] for i in edits if i.startswith(pre)],
                      key=lambda r: -r['start'])
        for r in todo:
            lead = r['raw'][:len(r['raw']) - len(r['raw'].lstrip())]
            tail = r['raw'][len(r['raw'].rstrip()):]
            new = esc_for(r.get('style', ''), r['new']) if fname == 'app.js' else r['new']
            src = src[:r['start']] + lead + new + tail + src[r['end']:]
            counts[fname] += 1
        if counts[fname]: open(P(fname), 'w', encoding='utf-8').write(src)

    # ---- data.js (and mirror into the build script) ----
    dedits = {i: edits[i] for i in edits if i.startswith('D')}
    if dedits:
        raw = open(P('data.js'), encoding='utf-8').read()
        payload = json.loads(raw[len('window.CAD='):-2])
        by_id = {d['id']: d for d in payload['datasets']}
        build = open(P('build_explorer_data.py'), encoding='utf-8').read()
        missed = []
        for i, e in dedits.items():
            path = manifest[i]['where']; old = manifest[i]['current']
            ds = path.split('.')[0]; node = by_id[ds]
            for part in re.findall(r'\w+|\[\d+\]', path[len(ds) + 1:]):
                key = int(part[1:-1]) if part.startswith('[') else part
                parent, last = node, key
                node = node[key]
            parent[last] = e['new']
            if old in build: build = build.replace(old, e['new'])
            else: missed.append(i)
            counts['data.js'] += 1
        open(P('data.js'), 'w', encoding='utf-8').write(
            'window.CAD=' + json.dumps(payload, separators=(',', ':')) + ';\n')
        open(P('build_explorer_data.py'), 'w', encoding='utf-8').write(build)
        if missed:
            print('NOTE: no matching literal in build_explorer_data.py for ' + ', '.join(missed))
            print('      data.js is updated; rerunning the build script would revert those.')

    for f, n in counts.items():
        if n: print(f'  {f:22} {n} string(s) updated')
    print(f'  {"TOTAL":22} {sum(counts.values())}')
    return 0

if __name__ == '__main__':
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else P('editing', 'CAD_Explorer_text.xlsx')))
